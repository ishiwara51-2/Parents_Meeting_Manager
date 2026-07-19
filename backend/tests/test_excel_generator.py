"""Excel生成サービスのテスト。

テスト対象: app.services.excel_generator.generate_excel

テストケース:
1. xlsxマジックバイト確認 -- 先頭が b"PK"（ZIP形式）
2. 3シート構成であること（日程案（出席番号）／番号-氏名対応表／日程案（氏名））
3. 「日程案（出席番号）」シートに出席番号が正しく配置されること
4. 「番号-氏名対応表」シートに全出席番号が列挙され、氏名列が空欄であること
5. 「日程案（氏名）」シートの対応セルに VLOOKUP 数式が入っていること
6. 「番号-氏名対応表」に氏名を記入した状態でVLOOKUP数式を評価すると
   氏名が正しく突合されること（openpyxlの数式パーサではなく、実際にopenpyxlで
   計算値を検証するのは困難なため、数式文字列に必要な要素が含まれることを確認する）
7. 未配置の生徒番号がフッターに表示されること
"""

from __future__ import annotations

import datetime
from io import BytesIO

from openpyxl import load_workbook

from app.models import Project, TimeSlot
from app.models.draft import Assignment, Draft
from app.services.excel_generator import (
    SHEET_ASSIGNMENTS,
    SHEET_ASSIGNMENTS_NAMED,
    SHEET_NAME_MAP,
    generate_excel,
)

_TZ_JST = datetime.timezone(datetime.timedelta(hours=9))


def _make_project(
    project_id: str = "test-project",
    display_name: str = "テストプロジェクト",
    candidate_dates: list[datetime.date] | None = None,
    candidate_time_slots: list[TimeSlot] | None = None,
    student_numbers: list[int] | None = None,
) -> Project:
    if candidate_dates is None:
        candidate_dates = [
            datetime.date(2026, 7, 15),
            datetime.date(2026, 7, 16),
        ]
    if candidate_time_slots is None:
        candidate_time_slots = [
            TimeSlot(start=datetime.time(16, 0), end=datetime.time(16, 20)),
            TimeSlot(start=datetime.time(16, 20), end=datetime.time(16, 40)),
        ]
    if student_numbers is None:
        student_numbers = [1, 2, 3]
    return Project(
        project_id=project_id,
        display_name=display_name,
        created_at=datetime.datetime(2026, 6, 24, 10, 0, 0, tzinfo=_TZ_JST),
        status="in_progress",
        slot_minutes=20,
        candidate_dates=candidate_dates,
        candidate_time_slots=candidate_time_slots,
        student_numbers=student_numbers,
    )


def _make_draft(
    project_id: str = "test-project",
    assignments: list[Assignment] | None = None,
    unassigned_students: list[int] | None = None,
) -> Draft:
    if assignments is None:
        assignments = [
            Assignment(
                student_number=1,
                date=datetime.date(2026, 7, 15),
                start=datetime.time(16, 0),
                end=datetime.time(16, 20),
            ),
        ]
    if unassigned_students is None:
        unassigned_students = []
    return Draft(
        project_id=project_id,
        saved_at=datetime.datetime(2026, 6, 24, 20, 0, 0, tzinfo=_TZ_JST),
        locked=True,
        assignments=assignments,
        unassigned_students=unassigned_students,
        violated_constraints=[],
    )


class TestGenerateExcel:
    """generate_excel のテスト群。"""

    def test_xlsx_magic_bytes(self) -> None:
        """生成されたxlsxがZIP形式のマジックバイト(PK)で始まること。"""
        project = _make_project()
        draft = _make_draft()
        result = generate_excel(project, draft)

        assert isinstance(result, bytes), "generate_excel は bytes を返すこと"
        assert result[:2] == b"PK", f"先頭2バイトが PK でない: {result[:2]!r}"

    def test_has_three_sheets_with_expected_names(self) -> None:
        """3シート（日程案（出席番号）／番号-氏名対応表／日程案（氏名））が生成されること。"""
        project = _make_project()
        draft = _make_draft()
        wb = load_workbook(BytesIO(generate_excel(project, draft)))

        assert wb.sheetnames == [
            SHEET_ASSIGNMENTS,
            SHEET_NAME_MAP,
            SHEET_ASSIGNMENTS_NAMED,
        ]

    def test_assignment_sheet_contains_student_numbers(self) -> None:
        """「日程案（出席番号）」シートに割り当て済みの出席番号が配置されること。"""
        project = _make_project(student_numbers=[1, 2])
        assignments = [
            Assignment(
                student_number=1,
                date=datetime.date(2026, 7, 15),
                start=datetime.time(16, 0),
                end=datetime.time(16, 20),
            ),
            Assignment(
                student_number=2,
                date=datetime.date(2026, 7, 16),
                start=datetime.time(16, 20),
                end=datetime.time(16, 40),
            ),
        ]
        draft = _make_draft(assignments=assignments)
        wb = load_workbook(BytesIO(generate_excel(project, draft)))
        ws = wb[SHEET_ASSIGNMENTS]

        assert ws["A1"].value == "時間帯"
        assert ws["B1"].value == "07/15"
        assert ws["C1"].value == "07/16"
        # 16:00-16:20 行 (row 2) の 07/15 列 (col B) に student_number=1
        assert ws["B2"].value == 1
        # 16:20-16:40 行 (row 3) の 07/16 列 (col C) に student_number=2
        assert ws["C3"].value == 2
        # 空きコマは None
        assert ws["C2"].value is None

    def test_name_map_sheet_lists_all_student_numbers_with_blank_names(self) -> None:
        """「番号-氏名対応表」シートに全出席番号が列挙され、氏名列が空欄であること。"""
        project = _make_project(student_numbers=[3, 1, 2])
        draft = _make_draft()
        wb = load_workbook(BytesIO(generate_excel(project, draft)))
        ws = wb[SHEET_NAME_MAP]

        assert ws["A1"].value == "出席番号"
        assert ws["B1"].value == "氏名"
        # ソート済みで出力される
        assert ws["A2"].value == 1
        assert ws["A3"].value == 2
        assert ws["A4"].value == 3
        assert ws["B2"].value is None
        assert ws["B3"].value is None
        assert ws["B4"].value is None

    def test_named_sheet_cells_contain_vlookup_formula(self) -> None:
        """「日程案（氏名）」シートの対応セルに番号-氏名対応表を参照するVLOOKUP数式が入ること。"""
        project = _make_project(student_numbers=[1])
        assignments = [
            Assignment(
                student_number=1,
                date=datetime.date(2026, 7, 15),
                start=datetime.time(16, 0),
                end=datetime.time(16, 20),
            ),
        ]
        draft = _make_draft(assignments=assignments)
        wb = load_workbook(BytesIO(generate_excel(project, draft)))
        ws = wb[SHEET_ASSIGNMENTS_NAMED]

        formula = ws["B2"].value
        assert isinstance(formula, str) and formula.startswith("=")
        assert f"'{SHEET_ASSIGNMENTS}'!B2" in formula
        assert f"'{SHEET_NAME_MAP}'!$A:$B" in formula
        assert "VLOOKUP" in formula

        # ヘッダ・時間帯ラベルは静的な値としてコピーされる
        assert ws["A1"].value == "時間帯"
        assert ws["B1"].value == "07/15"
        assert ws["A2"].value == "16:00-16:20"

    def test_unassigned_students_shown_in_footer(self) -> None:
        """未配置の出席番号が「日程案（出席番号）」シートのフッターに表示されること。"""
        project = _make_project(student_numbers=[1, 2, 3])
        draft = _make_draft(
            assignments=[
                Assignment(
                    student_number=1,
                    date=datetime.date(2026, 7, 15),
                    start=datetime.time(16, 0),
                    end=datetime.time(16, 20),
                ),
            ],
            unassigned_students=[3, 2],
        )
        wb = load_workbook(BytesIO(generate_excel(project, draft)))
        ws = wb[SHEET_ASSIGNMENTS]

        footer_texts = [
            str(v)
            for r in range(1, ws.max_row + 1)
            if (v := ws.cell(row=r, column=1).value) is not None
        ]
        assert any(
            "未配置" in text and "2" in text and "3" in text for text in footer_texts
        ), f"未配置フッターが見つからない: {footer_texts!r}"
